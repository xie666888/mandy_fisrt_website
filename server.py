#!/usr/bin/env python3
import base64
import hashlib
import hmac
import html
import io
import json
import logging
import mimetypes
import os
import secrets
import sqlite3
import threading
import time
import re
from http import cookies
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse

try:
    from PIL import Image as _PILImage
except ImportError:
    _PILImage = None

logger = logging.getLogger("luxe_catalog")
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)


ROOT = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("CATALOG_DB", ROOT / "catalog.sqlite3"))
UPLOAD_DIR = Path(os.environ.get("UPLOAD_DIR", ROOT / "src" / "assets" / "uploads"))
THUMB_DIR = UPLOAD_DIR / "_thumbs"
HOST = os.environ.get("HOST", "0.0.0.0")
PORT = int(os.environ.get("PORT", "4173"))
SESSION_COOKIE = "luxe_admin_session"
SESSION_TTL_SECONDS = 60 * 60 * 12
LOGIN_MAX_FAILURES_PER_USER = 5
LOGIN_MAX_FAILURES_PER_IP = 20
LOGIN_FAILURE_WINDOW_SECONDS = 10 * 60
LOGIN_FAILURE_DELAY_SECONDS = 0.35
MAX_IMAGE_UPLOAD_BYTES = 35 * 1024 * 1024
MAX_IMAGE_BATCH_UPLOAD_BYTES = 220 * 1024 * 1024
MAX_JSON_BYTES = 10 * 1024 * 1024
TRUST_PROXY = os.environ.get("TRUST_PROXY", "0") == "1"
HTTPS_ENABLED = os.environ.get("HTTPS", "0") == "1"
PUBLIC_BASE_URL = os.environ.get("PUBLIC_BASE_URL", "https://bebeauty.top").rstrip("/")

SEO_SITE_NAME = "BeBeauty Wholesale Catalog"
SEO_HOME_DESCRIPTION = (
    "Browse wholesale cosmetics and beauty products by brand, category, SKU, "
    "shade and price. Create an order number for supplier confirmation."
)
SEO_FAQ = [
    (
        "Are these products authentic?",
        "No. These are replica products and are not sold as authentic branded goods.",
    ),
    (
        "Can every product be scanned in the Sephora app?",
        "No. Sephora app scanning is not guaranteed for these products.",
    ),
    (
        "How does the order process work?",
        "Add products, shades and quantities to the cart, create an order number, "
        "then confirm availability and the Alibaba payment link through WhatsApp.",
    ),
    (
        "How long does delivery take?",
        "Goods are normally sent to the forwarder within three working days after "
        "payment. Forwarder delivery is usually 12 to 15 days.",
    ),
    (
        "What happens if goods are damaged during transportation?",
        "A replacement can be sent with the next order or the damaged item cost can "
        "be refunded after confirmation.",
    ),
]

_local = threading.local()


def connect():
    conn = getattr(_local, "conn", None)
    if conn is None:
        conn = sqlite3.connect(DB_PATH, timeout=30, check_same_thread=False)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA busy_timeout=30000")
        conn.execute("PRAGMA synchronous=NORMAL")
        conn.execute("PRAGMA cache_size=-64000")
        conn.execute("PRAGMA foreign_keys=ON")
        _local.conn = conn
    return conn


def now():
    return int(time.time())


def json_response(handler, status, payload):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Cache-Control", "no-store, max-age=0")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def text_response(handler, status, text, content_type, cache_control="public, max-age=3600"):
    body = text.encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", content_type)
    handler.send_header("Cache-Control", cache_control)
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def binary_response(handler, status, data, content_type, filename=None):
    handler.send_response(status)
    handler.send_header("Content-Type", content_type)
    handler.send_header("Content-Length", str(len(data)))
    if filename:
        handler.send_header("Content-Disposition", f'attachment; filename="{filename}"')
    handler.end_headers()
    handler.wfile.write(data)


def read_json(handler):
    return read_json_limited(handler, MAX_JSON_BYTES)


def read_json_limited(handler, max_bytes):
    length = int(handler.headers.get("Content-Length", "0"))
    if length > max_bytes:
        raise ValueError(f"File is too large. Maximum upload size is {max_bytes // 1024 // 1024}MB.")
    if not length:
        return {}
    raw = handler.rfile.read(length)
    return json.loads(raw.decode("utf-8"))


def parse_multipart_upload(handler, max_bytes):
    content_type = handler.headers.get("Content-Type", "")
    match = re.search(r'boundary=(?:"([^"]+)"|([^;]+))', content_type, re.I)
    if not match:
        raise ValueError("Multipart upload boundary is missing.")
    length = int(handler.headers.get("Content-Length", "0"))
    if length > max_bytes:
        raise ValueError(f"Upload is too large. Maximum batch size is {max_bytes // 1024 // 1024}MB.")
    raw = handler.rfile.read(length)
    boundary = ("--" + (match.group(1) or match.group(2)).strip()).encode("utf-8")
    files = []
    for part in raw.split(boundary):
        part = part.strip(b"\r\n")
        if not part or part == b"--":
            continue
        if part.endswith(b"--"):
            part = part[:-2].rstrip(b"\r\n")
        header_blob, separator, body = part.partition(b"\r\n\r\n")
        if not separator:
            continue
        headers = header_blob.decode("utf-8", errors="replace")
        disposition = next((line for line in headers.split("\r\n") if line.lower().startswith("content-disposition:")), "")
        if not disposition or "filename=" not in disposition:
            continue
        filename_match = re.search(r'filename="([^"]*)"|filename=([^;]+)', disposition)
        filename = (filename_match.group(1) or filename_match.group(2) if filename_match else "image").strip()
        if not filename:
            continue
        files.append({"filename": Path(filename).name, "data": body.rstrip(b"\r\n")})
    if not files:
        raise ValueError("No image files uploaded.")
    return files


def hash_password(password, salt=None):
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 120_000)
    return base64.b64encode(salt).decode(), base64.b64encode(digest).decode()


def verify_password(password, salt_b64, digest_b64):
    salt = base64.b64decode(salt_b64.encode())
    expected = base64.b64decode(digest_b64.encode())
    actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 120_000)
    return hmac.compare_digest(actual, expected)


def client_ip(handler):
    if TRUST_PROXY:
        forwarded = handler.headers.get("X-Forwarded-For", "").split(",")[0].strip()
        real_ip = handler.headers.get("X-Real-IP", "").strip()
        if forwarded or real_ip:
            return forwarded or real_ip
    return handler.client_address[0]


def login_key(username):
    return (str(username or "").strip().lower() or "__empty__")[:120]


def login_is_rate_limited(conn, ip, username):
    cutoff = now() - LOGIN_FAILURE_WINDOW_SECONDS
    conn.execute("DELETE FROM login_failures WHERE attempted_at < ?", (cutoff,))
    key = login_key(username)
    row = conn.execute(
        """
        SELECT
            SUM(CASE WHEN username = ? THEN 1 ELSE 0 END) AS pair_count,
            COUNT(*) AS ip_count
        FROM login_failures
        WHERE ip = ? AND attempted_at >= ?
        """,
        (key, ip, cutoff),
    ).fetchone()
    pair_count = int(row["pair_count"] or 0)
    ip_count = int(row["ip_count"] or 0)
    return pair_count >= LOGIN_MAX_FAILURES_PER_USER or ip_count >= LOGIN_MAX_FAILURES_PER_IP


def record_login_failure(conn, ip, username):
    conn.execute(
        "INSERT INTO login_failures (ip, username, attempted_at) VALUES (?, ?, ?)",
        (ip, login_key(username), now()),
    )


def clear_login_failures(conn, ip, username):
    conn.execute("DELETE FROM login_failures WHERE ip = ? AND username = ?", (ip, login_key(username)))


def product_from_row(row):
    product = dict(row)
    product["published"] = bool(product.get("published"))
    product["archived_at"] = int(product.get("archived_at") or 0)
    product["price"] = float(product.get("price") or 0)
    product["stock"] = int(product.get("stock") or 0)
    product["weight"] = float(product.get("weight") or 0)
    try:
        product["colors"] = json.loads(product.get("colors_json") or "[]")
    except json.JSONDecodeError:
        product["colors"] = []
    try:
        product["images"] = json.loads(product.get("images_json") or "[]")
    except json.JSONDecodeError:
        product["images"] = []
    if product.get("image") and product["image"] not in product["images"]:
        product["images"] = [product["image"], *product["images"]]
    product.pop("colors_json", None)
    product.pop("images_json", None)
    return product


def product_page_url(product_id):
    return f"{PUBLIC_BASE_URL}/product/{quote(str(product_id), safe='')}"


def absolute_media_url(value):
    value = str(value or "").strip().replace("\\", "/")
    if not value:
        return ""
    if value.startswith(("http://", "https://")):
        return value
    path = value.removeprefix("./").lstrip("/")
    return f"{PUBLIC_BASE_URL}/{quote(path, safe='/%:@?&=+$,;~.-_')}"


def seo_description(product):
    description = re.sub(r"\s+", " ", str(product.get("description") or "")).strip()
    identity = " ".join(
        part
        for part in [
            str(product.get("name") or "").strip(),
            f"by {str(product.get('brand') or '').strip()}" if product.get("brand") else "",
            f"SKU {str(product.get('sku') or '').strip()}" if product.get("sku") else "",
        ]
        if part
    )
    suffix = "Wholesale beauty catalog with shade selection and order-number confirmation."
    text = ". ".join(part.rstrip(".") for part in [identity, description, suffix] if part)
    return text[:157].rstrip(" ,.;") + "."


def seo_json(data):
    return json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")


def build_product_page(product, related_products):
    name = str(product.get("name") or product.get("sku") or "Wholesale beauty product")
    brand = str(product.get("brand") or "Unbranded")
    category = str(product.get("category") or "Cosmetics")
    sku = str(product.get("sku") or product.get("id") or "")
    description = seo_description(product)
    canonical = product_page_url(product.get("id") or sku)
    image = absolute_media_url(product.get("image"))
    images = [absolute_media_url(item) for item in product.get("images") or []]
    images = [item for item in dict.fromkeys([image, *images]) if item]
    title = f"{name} Wholesale | {brand} | BeBeauty"
    price = f"{float(product.get('price') or 0):.2f}"

    product_schema = {
        "@context": "https://schema.org",
        "@type": "Product",
        "name": name,
        "sku": sku,
        "category": category,
        "description": description,
        "url": canonical,
        "brand": {"@type": "Brand", "name": brand},
        "offers": {
            "@type": "Offer",
            "url": canonical,
            "priceCurrency": "USD",
            "price": price,
            "seller": {
                "@type": "Organization",
                "name": SEO_SITE_NAME,
                "url": PUBLIC_BASE_URL,
            },
        },
    }
    if images:
        product_schema["image"] = images
    breadcrumb_schema = {
        "@context": "https://schema.org",
        "@type": "BreadcrumbList",
        "itemListElement": [
            {
                "@type": "ListItem",
                "position": 1,
                "name": "Wholesale beauty products",
                "item": f"{PUBLIC_BASE_URL}/",
            },
            {
                "@type": "ListItem",
                "position": 2,
                "name": brand,
                "item": f"{PUBLIC_BASE_URL}/?brand={quote(brand)}",
            },
            {"@type": "ListItem", "position": 3, "name": name, "item": canonical},
        ],
    }
    faq_schema = {
        "@context": "https://schema.org",
        "@type": "FAQPage",
        "mainEntity": [
            {
                "@type": "Question",
                "name": question,
                "acceptedAnswer": {"@type": "Answer", "text": answer},
            }
            for question, answer in SEO_FAQ
        ],
    }
    related_html = "".join(
        f'<li><a href="{html.escape(product_page_url(item["id"]), quote=True)}">'
        f'{html.escape(str(item["name"]))}</a> by {html.escape(str(item["brand"]))}</li>'
        for item in related_products
    )
    image_html = (
        f'<img src="{html.escape(image, quote=True)}" alt="{html.escape(name, quote=True)}" '
        'width="395" height="395" loading="eager">'
        if image
        else ""
    )
    return f"""<!doctype html>
<html lang="en">
  <head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <base href="/">
    <title>{html.escape(title)}</title>
    <meta name="description" content="{html.escape(description, quote=True)}">
    <meta name="robots" content="index,follow,max-image-preview:large,max-snippet:-1">
    <link rel="canonical" href="{html.escape(canonical, quote=True)}">
    <link rel="alternate" hreflang="en" href="{html.escape(canonical, quote=True)}">
    <link rel="alternate" hreflang="x-default" href="{html.escape(canonical, quote=True)}">
    <meta property="og:type" content="product">
    <meta property="og:site_name" content="{html.escape(SEO_SITE_NAME, quote=True)}">
    <meta property="og:title" content="{html.escape(title, quote=True)}">
    <meta property="og:description" content="{html.escape(description, quote=True)}">
    <meta property="og:url" content="{html.escape(canonical, quote=True)}">
    {f'<meta property="og:image" content="{html.escape(image, quote=True)}">' if image else ""}
    <meta name="twitter:card" content="summary_large_image">
    <link rel="stylesheet" href="/src/styles.css?v=20260703-imgfix1">
    <script type="application/ld+json">{seo_json(product_schema)}</script>
    <script type="application/ld+json">{seo_json(breadcrumb_schema)}</script>
    <script type="application/ld+json">{seo_json(faq_schema)}</script>
  </head>
  <body>
    <div id="app">
      <main class="detail seo-product">
        <nav aria-label="Breadcrumb"><a href="/">Wholesale beauty products</a> / {html.escape(brand)}</nav>
        <article class="detail-layout">
          <div class="gallery"><div class="product-image">{image_html}</div></div>
          <section>
            <p>{html.escape(brand)} · {html.escape(category)} · SKU {html.escape(sku)}</p>
            <h1>{html.escape(name)}</h1>
            <p class="price">${price}</p>
            <h2>Product information</h2>
            <p>{html.escape(str(product.get("description") or description))}</p>
            <p><a href="/#products">Browse the wholesale catalog</a></p>
          </section>
        </article>
        <section><h2>Related wholesale products</h2><ul>{related_html}</ul></section>
      </main>
    </div>
    <script src="/src/app.js?v=20260703-imgfix1"></script>
  </body>
</html>
"""


def split_color_text(value):
    return [item.strip() for item in str(value or "").split("/") if item.strip()]


def normalize_product(data):
    sku = str(data.get("sku") or data.get("id") or "").strip()
    product_id = str(data.get("id") or sku).strip()
    name = str(data.get("name") or "").strip()
    if not product_id or not sku or not name:
        raise ValueError("SKU and Item Name are required.")
    colors = data.get("colors") or []
    if isinstance(colors, str):
        colors = split_color_text(colors)
    else:
        colors = [color for item in colors for color in split_color_text(item)]
    images = data.get("images") or []
    if isinstance(images, str):
        images = [item.strip() for item in images.split(",") if item.strip()]
    image = str(data.get("image") or "").strip()
    if image and image not in images:
        images = [image, *images]
    return {
        "id": product_id,
        "sku": sku,
        "brand": str(data.get("brand") or "UNBRANDED").strip(),
        "name": name,
        "category": str(data.get("category") or "Cosmetics").strip(),
        "price": float(data.get("price") or 0),
        "stock": int(float(data.get("stock") or 0)),
        "weight": float(data.get("weight") or 0),
        "description": str(data.get("description") or "").strip(),
        "colors_json": json.dumps(colors, ensure_ascii=False),
        "image": image or (images[0] if images else ""),
        "images_json": json.dumps(images, ensure_ascii=False),
        "published": 1 if data.get("published", True) else 0,
        "archived_at": int(float(data.get("archived_at") or 0)),
        "sort_order": int(float(data.get("sort_order") or 0)),
        "updated_at": now(),
    }


def unique_product_id(conn, preferred_id):
    base = str(preferred_id or "").strip()
    if not base:
        raise ValueError("SKU and Item Name are required.")
    if not conn.execute("SELECT 1 FROM products WHERE id = ?", (base,)).fetchone():
        return base
    suffix = 2
    while True:
        candidate = f"{base}--{suffix}"
        if not conn.execute("SELECT 1 FROM products WHERE id = ?", (candidate,)).fetchone():
            return candidate
        suffix += 1


def assign_unique_product_id(conn, data):
    sku = str(data.get("sku") or data.get("id") or "").strip()
    preferred = str(data.get("id") or sku).strip()
    data["id"] = unique_product_id(conn, preferred)
    if not data.get("sku"):
        data["sku"] = sku
    return data["id"]


def assign_unique_product_ids_batch(conn, products):
    existing = {row["id"] for row in conn.execute("SELECT id FROM products").fetchall()}
    used = set(existing)
    for product in products:
        sku = str(product.get("sku") or product.get("id") or "").strip()
        preferred = str(product.get("id") or sku).strip()
        if not preferred:
            raise ValueError("SKU and Item Name are required.")
        candidate = preferred
        suffix = 2
        while candidate in used:
            candidate = f"{preferred}--{suffix}"
            suffix += 1
        product["id"] = candidate
        if not product.get("sku"):
            product["sku"] = sku
        used.add(candidate)


def upsert_product(conn, data):
    product = normalize_product(data)
    conn.execute(
        """
        INSERT INTO products (
            id, sku, brand, name, category, price, stock, weight,
            description, colors_json, image, images_json, published, archived_at, sort_order, updated_at
        )
        VALUES (
            :id, :sku, :brand, :name, :category, :price, :stock, :weight,
            :description, :colors_json, :image, :images_json, :published, :archived_at, :sort_order, :updated_at
        )
        ON CONFLICT(id) DO UPDATE SET
            sku = excluded.sku,
            brand = excluded.brand,
            name = excluded.name,
            category = excluded.category,
            price = excluded.price,
            stock = excluded.stock,
            weight = excluded.weight,
            description = excluded.description,
            colors_json = excluded.colors_json,
            image = excluded.image,
            images_json = excluded.images_json,
            published = excluded.published,
            archived_at = excluded.archived_at,
            sort_order = excluded.sort_order,
            updated_at = excluded.updated_at
        """,
        product,
    )
    return product["id"]


def seed_products(conn):
    count = conn.execute("SELECT COUNT(*) AS count FROM products").fetchone()["count"]
    if count:
        return
    catalog_path = ROOT / "src" / "data" / "catalog.json"
    if not catalog_path.exists():
        return
    products = json.loads(catalog_path.read_text(encoding="utf-8"))
    assign_unique_product_ids_batch(conn, products)
    for product in products:
        upsert_product(conn, product)


def write_static_catalog(products):
    catalog_path = ROOT / "src" / "data" / "catalog.json"
    js_path = ROOT / "src" / "data" / "catalog.js"
    catalog_path.parent.mkdir(parents=True, exist_ok=True)
    catalog_path.write_text(json.dumps(products, ensure_ascii=False, indent=2), encoding="utf-8")
    js_path.write_text(
        "window.CATALOG_PRODUCTS = " + json.dumps(products, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )


def init_db():
    with connect() as conn:
        conn.execute("PRAGMA journal_mode=WAL")
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS products (
                id TEXT PRIMARY KEY,
                sku TEXT NOT NULL,
                brand TEXT NOT NULL,
                name TEXT NOT NULL,
                category TEXT NOT NULL,
                price REAL NOT NULL DEFAULT 0,
                stock INTEGER NOT NULL DEFAULT 0,
                weight REAL NOT NULL DEFAULT 0,
                description TEXT NOT NULL DEFAULT '',
                colors_json TEXT NOT NULL DEFAULT '[]',
                image TEXT NOT NULL DEFAULT '',
                images_json TEXT NOT NULL DEFAULT '[]',
                published INTEGER NOT NULL DEFAULT 1,
                archived_at INTEGER NOT NULL DEFAULT 0,
                sort_order INTEGER NOT NULL DEFAULT 0,
                updated_at INTEGER NOT NULL DEFAULT 0
            );

            CREATE TABLE IF NOT EXISTS admin_users (
                username TEXT PRIMARY KEY,
                password_salt TEXT NOT NULL,
                password_hash TEXT NOT NULL,
                updated_at INTEGER NOT NULL
            );

            CREATE TABLE IF NOT EXISTS sessions (
                token TEXT PRIMARY KEY,
                username TEXT NOT NULL,
                expires_at INTEGER NOT NULL
            );

            CREATE TABLE IF NOT EXISTS login_failures (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                ip TEXT NOT NULL,
                username TEXT NOT NULL,
                attempted_at INTEGER NOT NULL
            );

            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS orders (
                order_no TEXT PRIMARY KEY,
                country TEXT NOT NULL,
                items_json TEXT NOT NULL,
                product_total REAL NOT NULL DEFAULT 0,
                total_weight REAL NOT NULL DEFAULT 0,
                shipping REAL NOT NULL DEFAULT 0,
                total REAL NOT NULL DEFAULT 0,
                created_at INTEGER NOT NULL
            );

            CREATE INDEX IF NOT EXISTS idx_login_failures_ip_time
                ON login_failures (ip, attempted_at);

            CREATE INDEX IF NOT EXISTS idx_login_failures_ip_user_time
                ON login_failures (ip, username, attempted_at);

            CREATE INDEX IF NOT EXISTS idx_products_public_list
                ON products (published, archived_at, sort_order, brand, name);

            CREATE INDEX IF NOT EXISTS idx_products_admin_list
                ON products (archived_at, sort_order, brand, name);

            CREATE INDEX IF NOT EXISTS idx_products_brand_category
                ON products (brand, category);

            CREATE INDEX IF NOT EXISTS idx_sessions_expires_at
                ON sessions (expires_at);

            CREATE INDEX IF NOT EXISTS idx_orders_created_at
                ON orders (created_at);
            """
        )

        username = os.environ.get("ADMIN_USERNAME", "admin")
        password = os.environ.get("ADMIN_PASSWORD")
        exists = conn.execute("SELECT 1 FROM admin_users WHERE username = ?", (username,)).fetchone()
        if not exists:
            if not password:
                raise RuntimeError("ADMIN_PASSWORD environment variable is required for initial setup.")
            salt, digest = hash_password(password)
            conn.execute(
                "INSERT INTO admin_users (username, password_salt, password_hash, updated_at) VALUES (?, ?, ?, ?)",
                (username, salt, digest, now()),
            )

        conn.execute(
            "INSERT OR IGNORE INTO settings (key, value) VALUES (?, ?)",
            ("whatsapp", os.environ.get("WHATSAPP_NUMBER", "8613800000000")),
        )

        columns = {row["name"] for row in conn.execute("PRAGMA table_info(products)").fetchall()}
        if "images_json" not in columns:
            conn.execute("ALTER TABLE products ADD COLUMN images_json TEXT NOT NULL DEFAULT '[]'")
        if "sort_order" not in columns:
            conn.execute("ALTER TABLE products ADD COLUMN sort_order INTEGER NOT NULL DEFAULT 0")
        if "archived_at" not in columns:
            conn.execute("ALTER TABLE products ADD COLUMN archived_at INTEGER NOT NULL DEFAULT 0")
        seed_products(conn)
        conn.commit()


def parse_cookie(header):
    jar = cookies.SimpleCookie()
    if header:
        jar.load(header)
    return jar


def shipping_cost(country, weight_kg):
    weight = max(0, float(weight_kg or 0))
    if country == "United States":
        if weight <= 0.7:
            return 48.8
        if weight <= 1.2:
            return 65.4
        if weight <= 2:
            return 79.8
        if weight <= 3:
            return 84.2
        if weight <= 4:
            return 94.2
        if weight <= 5:
            return weight * 21.3
        if weight <= 15:
            return weight * 19.3
        return weight * 17.3

    if weight <= 0.7:
        return 29.8
    if weight <= 1.2:
        return 35.2
    if weight <= 2:
        return 43.8
    if weight <= 3:
        return 49.8
    if weight <= 4:
        return 59.5
    if weight <= 5:
        return 71.6
    if weight <= 15:
        return weight * 10.3
    return weight * 9.85


def local_image_path(image_url):
    if not image_url or "://" in image_url:
        return None
    safe = Path(str(image_url).replace("\\", "/").lstrip("./").lstrip("/"))
    if ".." in safe.parts:
        return None
    if safe.parts[:3] == ("src", "assets", "uploads"):
        uploaded = UPLOAD_DIR / Path(*safe.parts[3:])
        if uploaded.exists() and uploaded.is_file():
            return uploaded
    path = ROOT / safe
    if path.exists() and path.is_file():
        return path
    return None


def thumbnail_path_for(image_url, size):
    source = local_image_path(image_url)
    if not source:
        return None
    try:
        mtime = int(source.stat().st_mtime)
    except OSError:
        return None
    key = hashlib.sha256(f"{image_url}|{size}|{mtime}".encode("utf-8")).hexdigest()[:24]
    return source, THUMB_DIR / f"{key}-{size}.jpg"


def make_thumbnail(image_url, size=160):
    result = thumbnail_path_for(image_url, size)
    if not result:
        return None
    source, thumb_path = result
    if thumb_path.exists():
        return thumb_path
    if _PILImage is None:
        return source
    THUMB_DIR.mkdir(parents=True, exist_ok=True)
    try:
        with _PILImage.open(source) as image:
            image.thumbnail((size, size), _PILImage.Resampling.LANCZOS)
            if image.mode in {"RGBA", "LA", "P"}:
                image = image.convert("RGBA")
                background = _PILImage.new("RGB", image.size, (255, 255, 255))
                background.paste(image, mask=image.getchannel("A"))
                image = background
            else:
                image = image.convert("RGB")
            image.save(thumb_path, "JPEG", quality=72, optimize=True, progressive=True)
        return thumb_path
    except Exception:
        logger.warning("Thumbnail generation failed for %s", image_url, exc_info=True)
        return source


def safe_file_stem(value):
    keep = []
    for char in str(value or "image"):
        if char.isascii() and (char.isalnum() or char in {"-", "_", "."}):
            keep.append(char)
        else:
            keep.append("_")
    stem = "".join(keep).strip("._-")[:80]
    return stem or "image"


def image_extension(data, filename=""):
    lower = str(filename or "").lower()
    if data.startswith(b"\x89PNG"):
        return ".png"
    if data.startswith(b"\xff\xd8"):
        return ".jpg"
    if data.startswith(b"GIF"):
        return ".gif"
    if data.startswith(b"RIFF") and b"WEBP" in data[:16]:
        return ".webp"
    if lower.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp")):
        return ".jpg" if lower.endswith(".jpeg") else Path(lower).suffix
    raise ValueError("Only PNG, JPG, GIF, and WEBP images are supported.")


def save_uploaded_images(files):
    if not isinstance(files, list) or not files:
        raise ValueError("No image files uploaded.")
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    urls = []
    for item in files:
        filename = str(item.get("filename") or "image").strip()
        if item.get("data") is not None:
            raw = item["data"]
        else:
            content = str(item.get("contentBase64") or "")
            if not content:
                continue
            try:
                raw = base64.b64decode(content, validate=True)
            except Exception as exc:
                raise ValueError(f"{filename} is not a valid uploaded file.") from exc
        if len(raw) > MAX_IMAGE_UPLOAD_BYTES:
            raise ValueError(f"{filename} is too large. Maximum image size is {MAX_IMAGE_UPLOAD_BYTES // 1024 // 1024}MB.")
        ext = image_extension(raw, filename)
        stem = safe_file_stem(Path(filename).stem)
        file_name = f"{int(time.time())}-{secrets.token_hex(5)}-{stem}{ext}"
        target = UPLOAD_DIR / file_name
        target.write_bytes(raw)
        urls.append(f"./src/assets/uploads/{file_name}")
    if not urls:
        raise ValueError("No valid image files uploaded.")
    return urls


VALID_SHIPPING_COUNTRIES = {"Europe", "United States"}


def normalize_order_items(cart_items):
    requested = {}
    for item in cart_items:
        product_id = str(item.get("id") or item.get("sku") or "").strip()
        sku = str(item.get("sku") or product_id).strip()
        if not product_id:
            continue
        qty = max(1, int(float(item.get("qty") or 1)))
        color = str(item.get("color") or "Default").strip() or "Default"
        key = (product_id, sku, color)
        requested[key] = requested.get(key, 0) + qty
    if not requested:
        raise ValueError("Cart is empty.")
    return [{"id": product_id, "sku": sku, "color": color, "qty": qty} for (product_id, sku, color), qty in requested.items()]


def products_by_id(product_ids):
    unique_ids = sorted({product_id for product_id in product_ids if product_id})
    if not unique_ids:
        return {}
    placeholders = ",".join("?" for _ in unique_ids)
    with connect() as conn:
        rows = conn.execute(f"SELECT * FROM products WHERE id IN ({placeholders})", unique_ids).fetchall()
    return {row["id"]: product_from_row(row) for row in rows}


def prepare_order_data(cart_items, country):
    country = str(country or "Europe").strip()
    if country not in VALID_SHIPPING_COUNTRIES:
        raise ValueError("Shipping country must be Europe or United States.")

    order_items = normalize_order_items(cart_items)
    product_map = products_by_id([item["id"] for item in order_items])
    lines = []
    product_total = 0
    total_weight = 0
    for item in order_items:
        product_id = item["id"]
        product = product_map.get(product_id)
        if not product:
            raise ValueError(f"Product not found: {product_id}")
        sku = product.get("sku") or item["sku"] or product_id
        qty = item["qty"]
        unit_price = float(product.get("price") or 0)
        unit_weight = float(product.get("weight") or 0)
        extended = unit_price * qty
        line_weight = unit_weight * qty
        product_total += extended
        total_weight += line_weight
        lines.append(
            {
                "id": product.get("id") or product_id,
                "sku": sku,
                "brand": product.get("brand") or "",
                "name": product.get("name") or "",
                "color": item["color"],
                "qty": qty,
                "price": unit_price,
                "weight": unit_weight,
                "line_weight": line_weight,
                "extended": extended,
                "image": product.get("image") or "",
            }
        )

    shipping = round(shipping_cost(country, total_weight), 2)
    total = round(product_total + shipping, 2)
    return {
        "country": country,
        "items": lines,
        "product_total": round(product_total, 2),
        "total_weight": round(total_weight, 3),
        "shipping": shipping,
        "total": total,
    }


def generate_order_no():
    return "DL" + time.strftime("%Y%m%d%H%M%S", time.localtime()) + secrets.token_hex(2).upper()


def create_order_record(cart_items, country):
    order = prepare_order_data(cart_items, country)
    created_at = now()
    with connect() as conn:
        for _ in range(5):
            order_no = generate_order_no()
            try:
                conn.execute(
                    """
                    INSERT INTO orders (
                        order_no, country, items_json, product_total, total_weight,
                        shipping, total, created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        order_no,
                        order["country"],
                        json.dumps(order["items"], ensure_ascii=False),
                        order["product_total"],
                        order["total_weight"],
                        order["shipping"],
                        order["total"],
                        created_at,
                    ),
                )
                conn.commit()
                order["order_no"] = order_no
                order["created_at"] = created_at
                return order
            except sqlite3.IntegrityError:
                continue
    raise ValueError("Unable to generate a unique order number. Please try again.")


def load_order_record(order_no):
    order_no = str(order_no or "").strip()
    if not order_no or any(char not in "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789_-" for char in order_no):
        raise ValueError("Please enter a valid order number.")
    with connect() as conn:
        row = conn.execute("SELECT * FROM orders WHERE order_no = ?", (order_no,)).fetchone()
    if not row:
        raise ValueError(f"Order not found: {order_no}")
    order = dict(row)
    try:
        order["items"] = json.loads(order.pop("items_json") or "[]")
    except json.JSONDecodeError:
        order["items"] = []
    return order


def build_order_workbook_from_data(order):
    from copy import copy

    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Order"
    default_font = Font(name="Arial", size=11, family=2)
    wb._fonts[0] = default_font
    wb._named_styles["Normal"].font = default_font
    if order.get("order_no"):
        ws.append(["Order No", order.get("order_no")])
        ws.append(["Created at", time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(int(order.get("created_at") or now())))])
        ws.append([])

    headers = [
        "NO",
        "SKU",
        "Brand",
        "Item Name",
        "Picture",
        "Shade/Type",
        "Unit Price",
        "QTY",
        "Extended price",
        "Weight(kg)",
    ]
    ws.append(headers)
    header_row = ws.max_row
    header_fill = PatternFill("solid", fgColor="1F1A17")
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", family=2, color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for index, item in enumerate(order.get("items") or [], 1):
        unit_price = float(item.get("price") or 0)
        qty = max(1, int(float(item.get("qty") or 1)))
        unit_weight = float(item.get("weight") or 0)
        extended = float(item.get("extended") or unit_price * qty)
        line_weight = float(item.get("line_weight") or unit_weight * qty)
        row_number = ws.max_row + 1
        ws.append(
            [
                index,
                item.get("sku") or "",
                item.get("brand") or "",
                item.get("name") or "",
                "",
                item.get("color") or "Default",
                unit_price,
                qty,
                extended,
                line_weight,
            ]
        )
        ws.row_dimensions[row_number].height = 58
        image_path = local_image_path(item.get("image"))
        if image_path:
            try:
                image = XLImage(str(image_path))
                max_width, max_height = 72, 58
                ratio = min(max_width / image.width, max_height / image.height, 1)
                image.width = int(image.width * ratio)
                image.height = int(image.height * ratio)
                ws.add_image(image, f"E{row_number}")
            except Exception:
                pass

    summary_start = ws.max_row + 2
    summary_rows = [
        ("Shipping country", order.get("country") or "Europe"),
        ("Total product cost", float(order.get("product_total") or 0)),
        ("Total weight(kg)", float(order.get("total_weight") or 0)),
        ("SHIPPING COST", float(order.get("shipping") or 0)),
        ("TOTAL", float(order.get("total") or 0)),
    ]
    for label, value in summary_rows:
        ws.append(["", "", "", "", "", "", "", label, value, ""])

    widths = {
        "A": 8,
        "B": 18,
        "C": 18,
        "D": 42,
        "E": 14,
        "F": 18,
        "G": 12,
        "H": 10,
        "I": 16,
        "J": 14,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in range(header_row + 1, ws.max_row + 1):
        ws[f"G{row}"].number_format = "$0.00"
        ws[f"I{row}"].number_format = "$0.00"
        ws[f"J{row}"].number_format = "0.000"
    for row in range(summary_start, ws.max_row + 1):
        ws[f"H{row}"].font = Font(name="Arial", family=2, bold=True)
        ws[f"I{row}"].font = Font(name="Arial", family=2, bold=True)
        if ws[f"H{row}"].value in {"Total product cost", "SHIPPING COST", "TOTAL"}:
            ws[f"I{row}"].number_format = "$0.00"
    for row in ws.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name = "Arial"
            font.family = 2
            font.scheme = None
            font.charset = None
            cell.font = font
    ws.freeze_panes = f"A{header_row + 1}"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def build_order_workbook(cart_items, country):
    return build_order_workbook_from_data(prepare_order_data(cart_items, country))


class CatalogHandler(BaseHTTPRequestHandler):
    server_version = "LuxeCatalog/1.0"
    timeout = 120
    rbufsize = -1

    def end_headers(self):
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("X-Frame-Options", "DENY")
        self.send_header("Referrer-Policy", "same-origin")
        self.send_header("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
        if urlparse(self.path).path.startswith("/api/"):
            self.send_header("X-Robots-Tag", "noindex, nofollow")
        self.send_header(
            "Content-Security-Policy",
            "default-src 'self'; "
            "script-src 'self' https://cdn.sheetjs.com; "
            "style-src 'self' 'unsafe-inline'; "
            "img-src 'self' data: blob:; "
            "connect-src 'self'; "
            "frame-ancestors 'none'; "
            "base-uri 'self'; "
            "form-action 'self'",
        )
        super().end_headers()

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/robots.txt":
            return self.handle_robots()
        if parsed.path == "/sitemap.xml":
            return self.handle_sitemap()
        if parsed.path == "/llms.txt":
            return self.handle_llms()
        if parsed.path.startswith("/product/"):
            product_id = unquote(parsed.path.removeprefix("/product/")).rstrip("/")
            return self.handle_product_page(product_id)
        if parsed.path == "/api/products":
            return self.handle_products(parse_qs(parsed.query))
        if parsed.path == "/api/me":
            return self.handle_me()
        if parsed.path == "/api/settings":
            return self.handle_settings_get()
        if parsed.path == "/api/thumb":
            return self.handle_thumbnail(parse_qs(parsed.query))
        if parsed.path.startswith("/api/orders/") and parsed.path.endswith("/excel"):
            parts = parsed.path.strip("/").split("/")
            if len(parts) == 4 and parts[0] == "api" and parts[1] == "orders" and parts[3] == "excel":
                return self.require_admin(lambda: self.handle_order_download(parts[2]))
        if parsed.path.startswith("/api/"):
            return json_response(self, 404, {"error": "Not found"})
        return self.serve_static(parsed.path)

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/login":
            return self.handle_login()
        if parsed.path == "/api/logout":
            return self.handle_logout()
        if parsed.path == "/api/orders":
            return self.handle_order_create()
        if parsed.path == "/api/orders/excel":
            return self.handle_order_excel()
        if parsed.path == "/api/products":
            return self.require_admin(self.handle_product_create)
        if parsed.path == "/api/uploads/images":
            return self.require_admin(self.handle_image_upload)
        return json_response(self, 404, {"error": "Not found"})

    def do_PUT(self):
        parsed = urlparse(self.path)
        if parsed.path.startswith("/api/products/"):
            product_id = unquote(parsed.path.removeprefix("/api/products/"))
            return self.require_admin(lambda: self.handle_product_update(product_id))
        if parsed.path == "/api/settings":
            return self.require_admin(self.handle_settings_put)
        return json_response(self, 404, {"error": "Not found"})

    def do_DELETE(self):
        parsed = urlparse(self.path)
        if parsed.path.startswith("/api/products/"):
            product_id = unquote(parsed.path.removeprefix("/api/products/"))
            return self.require_admin(lambda: self.handle_product_delete(product_id))
        return json_response(self, 404, {"error": "Not found"})

    def current_user(self):
        if hasattr(self, "_cached_user"):
            return self._cached_user
        jar = parse_cookie(self.headers.get("Cookie"))
        morsel = jar.get(SESSION_COOKIE)
        user = None
        if morsel:
            token = morsel.value
            with connect() as conn:
                row = conn.execute("SELECT username, expires_at FROM sessions WHERE token = ?", (token,)).fetchone()
                if row and row["expires_at"] >= now():
                    user = row["username"]
                elif row:
                    conn.execute("DELETE FROM sessions WHERE token = ?", (token,))
                    conn.commit()
        self._cached_user = user
        return user

    def require_admin(self, action):
        if not self.current_user():
            return json_response(self, 401, {"error": "Authentication required"})
        return action()

    def handle_me(self):
        username = self.current_user()
        return json_response(self, 200, {"authenticated": bool(username), "username": username})

    def handle_login(self):
        try:
            data = read_json_limited(self, 4096)
        except Exception:
            return json_response(self, 400, {"error": "Invalid JSON"})
        username = str(data.get("username") or "").strip()
        password = str(data.get("password") or "")
        ip = client_ip(self)
        with connect() as conn:
            if login_is_rate_limited(conn, ip, username):
                conn.commit()
                return json_response(self, 429, {"error": "Too many failed login attempts. Please try again later."})
            row = conn.execute("SELECT * FROM admin_users WHERE username = ?", (username,)).fetchone()
            if not row or not verify_password(password, row["password_salt"], row["password_hash"]):
                record_login_failure(conn, ip, username)
                conn.commit()
                time.sleep(LOGIN_FAILURE_DELAY_SECONDS)
                return json_response(self, 401, {"error": "Username or password is incorrect."})
            token = secrets.token_urlsafe(32)
            expires_at = now() + SESSION_TTL_SECONDS
            conn.execute("INSERT INTO sessions (token, username, expires_at) VALUES (?, ?, ?)", (token, username, expires_at))
            clear_login_failures(conn, ip, username)
            conn.commit()
        body = json.dumps({"ok": True, "username": username}).encode("utf-8")
        cookie_flags = f"HttpOnly; SameSite=Lax; Path=/; Max-Age={SESSION_TTL_SECONDS}"
        if HTTPS_ENABLED:
            cookie_flags += "; Secure"
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store, max-age=0")
        self.send_header("Set-Cookie", f"{SESSION_COOKIE}={token}; {cookie_flags}")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_logout(self):
        jar = parse_cookie(self.headers.get("Cookie"))
        token = jar.get(SESSION_COOKIE).value if jar.get(SESSION_COOKIE) else ""
        with connect() as conn:
            if token:
                conn.execute("DELETE FROM sessions WHERE token = ?", (token,))
                conn.commit()
        body = json.dumps({"ok": True}).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store, max-age=0")
        self.send_header("Set-Cookie", f"{SESSION_COOKIE}=; HttpOnly; SameSite=Lax; Path=/; Max-Age=0")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_robots(self):
        content = (
            "User-agent: *\n"
            "Allow: /\n"
            "Disallow: /api/\n"
            f"Sitemap: {PUBLIC_BASE_URL}/sitemap.xml\n"
            f"Host: {urlparse(PUBLIC_BASE_URL).netloc}\n"
        )
        return text_response(self, 200, content, "text/plain; charset=utf-8")

    def handle_sitemap(self):
        with connect() as conn:
            rows = conn.execute(
                """
                SELECT id, name, image, updated_at
                FROM products
                WHERE published = 1 AND archived_at = 0
                ORDER BY updated_at DESC, id
                """
            ).fetchall()
        entries = [
            (
                f"  <url><loc>{html.escape(PUBLIC_BASE_URL + '/')}</loc>"
                "<changefreq>daily</changefreq><priority>1.0</priority></url>"
            )
        ]
        for row in rows:
            updated_at = int(row["updated_at"] or 0)
            lastmod = time.strftime("%Y-%m-%d", time.gmtime(updated_at)) if updated_at else ""
            image = absolute_media_url(row["image"])
            parts = [
                "  <url>",
                f"<loc>{html.escape(product_page_url(row['id']))}</loc>",
                f"<lastmod>{lastmod}</lastmod>" if lastmod else "",
                "<changefreq>weekly</changefreq><priority>0.8</priority>",
            ]
            if image:
                parts.extend(
                    [
                        "<image:image>",
                        f"<image:loc>{html.escape(image)}</image:loc>",
                        f"<image:title>{html.escape(str(row['name']))}</image:title>",
                        "</image:image>",
                    ]
                )
            parts.append("</url>")
            entries.append("".join(parts))
        sitemap = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9" '
            'xmlns:image="http://www.google.com/schemas/sitemap-image/1.1">\n'
            + "\n".join(entries)
            + "\n</urlset>\n"
        )
        return text_response(self, 200, sitemap, "application/xml; charset=utf-8")

    def handle_llms(self):
        with connect() as conn:
            product_count = conn.execute(
                "SELECT COUNT(*) FROM products WHERE published = 1 AND archived_at = 0"
            ).fetchone()[0]
            brands = conn.execute(
                """
                SELECT brand, COUNT(*) AS count
                FROM products
                WHERE published = 1 AND archived_at = 0 AND brand <> ''
                GROUP BY brand
                ORDER BY count DESC, brand
                LIMIT 30
                """
            ).fetchall()
            categories = conn.execute(
                """
                SELECT category, COUNT(*) AS count
                FROM products
                WHERE published = 1 AND archived_at = 0 AND category <> ''
                GROUP BY category
                ORDER BY count DESC, category
                """
            ).fetchall()
        brand_text = ", ".join(f"{row['brand']} ({row['count']})" for row in brands)
        category_text = ", ".join(f"{row['category']} ({row['count']})" for row in categories)
        content = f"""# {SEO_SITE_NAME}

> A B2B wholesale cosmetics catalog for browsing products, choosing shades and quantities, and creating an order number for supplier confirmation.

## Canonical website

- Website: {PUBLIC_BASE_URL}/
- Product sitemap: {PUBLIC_BASE_URL}/sitemap.xml
- Published products: {product_count}

## Catalog

- Major brands: {brand_text}
- Categories: {category_text}
- Product pages include SKU, brand, category, USD price, shade choices, product information and images.

## Ordering

1. Add products, shades and quantities to the cart.
2. Select Europe or United States for shipping.
3. Create an order number.
4. Send the order number through WhatsApp for availability confirmation.
5. The supplier provides an Alibaba payment link after details are agreed.

## Important product information

- Availability is confirmed manually after an order number is created.
- Sephora app scanning is not guaranteed.
- Products are replica products and are not sold as authentic branded goods.
- Goods are normally sent to the forwarder within three working days after payment.
- Forwarder delivery is usually 12 to 15 days.

## Contact workflow

Use the cart on {PUBLIC_BASE_URL}/ to create an order number, then continue the supplier conversation through the WhatsApp action provided by the website.
"""
        return text_response(self, 200, content, "text/plain; charset=utf-8")

    def handle_product_page(self, product_id):
        if not product_id:
            return self.send_error(404)
        with connect() as conn:
            row = conn.execute(
                """
                SELECT * FROM products
                WHERE id = ? AND published = 1 AND archived_at = 0
                """,
                (product_id,),
            ).fetchone()
            if not row:
                return self.send_error(404)
            product = product_from_row(row)
            related_rows = conn.execute(
                """
                SELECT * FROM products
                WHERE id <> ? AND published = 1 AND archived_at = 0
                    AND (brand = ? OR category = ?)
                ORDER BY CASE WHEN brand = ? THEN 0 ELSE 1 END,
                         sort_order ASC, name ASC
                LIMIT 8
                """,
                (product_id, product["brand"], product["category"], product["brand"]),
            ).fetchall()
        page = build_product_page(product, [product_from_row(item) for item in related_rows])
        return text_response(
            self,
            200,
            page,
            "text/html; charset=utf-8",
            cache_control="public, max-age=300",
        )

    def handle_products(self, query):
        include_drafts = query.get("include_drafts", ["0"])[0] == "1" and bool(self.current_user())
        where = "WHERE archived_at = 0" if include_drafts else "WHERE published = 1 AND archived_at = 0"
        with connect() as conn:
            rows = conn.execute(f"SELECT * FROM products {where} ORDER BY sort_order ASC, brand, name").fetchall()
        return json_response(self, 200, {"products": [product_from_row(row) for row in rows]})

    def handle_product_create(self):
        try:
            data = read_json(self)
            with connect() as conn:
                assign_unique_product_id(conn, data)
                product_id = upsert_product(conn, data)
                conn.commit()
            return json_response(self, 200, {"ok": True, "id": product_id})
        except ValueError as exc:
            return json_response(self, 400, {"error": str(exc)})

    def handle_product_update(self, product_id):
        try:
            data = read_json(self)
            data["id"] = product_id
            with connect() as conn:
                upsert_product(conn, data)
                conn.commit()
            return json_response(self, 200, {"ok": True, "id": product_id})
        except ValueError as exc:
            return json_response(self, 400, {"error": str(exc)})

    def handle_product_delete(self, product_id):
        with connect() as conn:
            cursor = conn.execute(
                "UPDATE products SET archived_at = ?, published = 0, updated_at = ? WHERE id = ?",
                (now(), now(), product_id),
            )
            conn.commit()
        if cursor.rowcount == 0:
            return json_response(self, 404, {"error": "Product not found"})
        return json_response(self, 200, {"ok": True, "archived": True})

    def handle_image_upload(self):
        try:
            content_type = self.headers.get("Content-Type", "")
            if content_type.lower().startswith("multipart/form-data"):
                files = parse_multipart_upload(self, MAX_IMAGE_BATCH_UPLOAD_BYTES)
            else:
                data = read_json_limited(self, MAX_IMAGE_BATCH_UPLOAD_BYTES * 2)
                files = data.get("files") or []
            urls = save_uploaded_images(files)
            return json_response(self, 200, {"ok": True, "urls": urls})
        except ValueError as exc:
            return json_response(self, 400, {"error": str(exc)})
        except Exception:
            logger.exception("Image upload failed")
            return json_response(self, 500, {"error": "Image upload failed. Please try again."})

    def handle_settings_get(self):
        with connect() as conn:
            rows = conn.execute("SELECT key, value FROM settings").fetchall()
        return json_response(self, 200, {"settings": {row["key"]: row["value"] for row in rows}})

    def handle_settings_put(self):
        data = read_json(self)
        whatsapp = str(data.get("whatsapp") or "").strip()
        with connect() as conn:
            conn.execute("INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value", ("whatsapp", whatsapp))
            conn.commit()
        return json_response(self, 200, {"ok": True})

    def handle_thumbnail(self, query):
        src = query.get("src", [""])[0]
        try:
            size = max(48, min(320, int(query.get("size", ["160"])[0] or 160)))
        except ValueError:
            size = 160
        thumb_path = make_thumbnail(src, size)
        if not thumb_path or not thumb_path.exists():
            return self.send_error(404)
        content_type = mimetypes.guess_type(thumb_path.name)[0] or "image/jpeg"
        file_size = thumb_path.stat().st_size
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Cache-Control", "public, max-age=31536000, immutable")
        self.send_header("Content-Length", str(file_size))
        self.end_headers()
        with thumb_path.open("rb") as f:
            while chunk := f.read(65536):
                self.wfile.write(chunk)

    def handle_order_create(self):
        try:
            data = read_json(self)
            items = data.get("items") or []
            if not isinstance(items, list) or not items:
                raise ValueError("Cart is empty.")
            order = create_order_record(items, data.get("country") or "Europe")
            return json_response(
                self,
                200,
                {
                    "ok": True,
                    "orderNo": order["order_no"],
                    "country": order["country"],
                    "productTotal": order["product_total"],
                    "totalWeight": order["total_weight"],
                    "shipping": order["shipping"],
                    "total": order["total"],
                },
            )
        except ValueError as exc:
            return json_response(self, 400, {"error": str(exc)})
        except Exception:
            logger.exception("Order creation failed")
            return json_response(self, 500, {"error": "Unable to create order. Please try again."})

    def handle_order_download(self, order_no):
        try:
            order = load_order_record(order_no)
            workbook_data = build_order_workbook_from_data(order)
            filename = f"order-{order['order_no']}.xlsx"
            return binary_response(
                self,
                200,
                workbook_data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                filename,
            )
        except ValueError as exc:
            return json_response(self, 404, {"error": str(exc)})
        except Exception:
            logger.exception("Order Excel download failed")
            return json_response(self, 500, {"error": "Unable to download order Excel. Please try again."})

    def handle_order_excel(self):
        try:
            data = read_json(self)
            items = data.get("items") or []
            country = str(data.get("country") or "Europe").strip()
            if country not in VALID_SHIPPING_COUNTRIES:
                raise ValueError("Shipping country must be Europe or United States.")
            if not isinstance(items, list) or not items:
                raise ValueError("Cart is empty.")
            workbook_data = build_order_workbook(items, country)
            filename = f"order-{int(time.time())}.xlsx"
            return binary_response(
                self,
                200,
                workbook_data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                filename,
            )
        except ValueError as exc:
            return json_response(self, 400, {"error": str(exc)})
        except Exception:
            logger.exception("Order Excel generation failed")
            return json_response(self, 500, {"error": "Unable to create order Excel. Please try again."})

    def serve_static(self, url_path):
        if url_path == "/":
            url_path = "/index.html"
        safe = Path(url_path.lstrip("/"))
        if ".." in safe.parts:
            return self.send_error(403)
        blocked_suffixes = {".py", ".sqlite3", ".db", ".zip", ".xlsx", ".xlsm", ".env", ".service"}
        blocked_dirs = {"scripts", "__pycache__"}
        if safe.suffix.lower() in blocked_suffixes or any(part in blocked_dirs or part.startswith(".") for part in safe.parts):
            return self.send_error(403)
        path = ROOT / safe
        if safe.parts[:3] == ("src", "assets", "uploads"):
            upload_path = UPLOAD_DIR / Path(*safe.parts[3:])
            if upload_path.exists() and upload_path.is_file():
                path = upload_path
        if not path.exists() or path.is_dir():
            return self.send_error(404)
        content_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        try:
            stat = path.stat()
        except OSError:
            return self.send_error(404)
        etag = f'"{stat.st_mtime_ns:x}-{stat.st_size:x}"'
        if self.headers.get("If-None-Match") == etag:
            self.send_response(304)
            self.end_headers()
            return
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(stat.st_size))
        self.send_header("ETag", etag)
        if path.suffix.lower() in {".html", ".js", ".css", ".json"}:
            self.send_header("Cache-Control", "no-store, max-age=0")
        elif content_type.startswith("image/"):
            self.send_header("Cache-Control", "public, max-age=31536000, immutable")
        self.end_headers()
        with path.open("rb") as f:
            while chunk := f.read(65536):
                self.wfile.write(chunk)

    def log_message(self, fmt, *args):
        logger.info("%s %s", self.address_string(), fmt % args)


if __name__ == "__main__":
    init_db()
    httpd = ThreadingHTTPServer((HOST, PORT), CatalogHandler)
    httpd.daemon_threads = True
    logger.info("Luxe catalog server running on http://%s:%s", HOST, PORT)
    logger.info("SQLite database: %s", DB_PATH)
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        logger.info("Server shutting down")
        httpd.shutdown()

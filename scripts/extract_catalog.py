import json
import re
import sys
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "2026.4.6 catalog_min.xlsx"
OUT = ROOT / "src" / "data" / "catalog.json"
JS_OUT = ROOT / "src" / "data" / "catalog.js"
IMAGE_DIR = ROOT / "src" / "assets" / "products"


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text.replace(".", "", 1).isdigit():
        text = text[:-2]
    return re.sub(r"\s+", " ", text)


def number(value, default=0):
    try:
        if value is None or value == "":
            return default
        return float(value)
    except Exception:
        return default


def category(name, info):
    text = f"{name} {info}".lower()
    pairs = [
        ("lip", "Lip"),
        ("rouge", "Lip"),
        ("gloss", "Lip"),
        ("blush", "Blush"),
        ("cream", "Skincare"),
        ("hand", "Skincare"),
        ("powder", "Face"),
        ("foundation", "Face"),
        ("concealer", "Face"),
        ("highlighter", "Face"),
        ("liner", "Eye"),
        ("mascara", "Eye"),
        ("eye", "Eye"),
        ("palette", "Palette"),
        ("perfume", "Fragrance"),
    ]
    for key, label in pairs:
        if key in text:
            return label
    return "Cosmetics"


def clean_color_list(items):
    cleaned = []
    for part in items:
        item = clean(part)
        item = re.sub(r"^colou?rs?\s*[:：]\s*", "", item, flags=re.I)
        item = item.strip("#/- ")
        if not item:
            continue
        if re.search(r"(?:size|gross|net|weight|cm|\*)", item, re.I):
            continue
        if re.match(r"^\d+(?:\.\d+)?g$", item, re.I):
            continue
        if item not in cleaned:
            cleaned.append(item)
    return cleaned[:36]


def colors(info, fallback=""):
    text = clean(info or fallback)
    match = re.search(r"(?:including\s+)?colou?rs?\s*[:：]\s*(.*)", text, re.I)
    if not match:
        return clean_color_list(str(fallback or "").split("/"))
    raw = match.group(1)
    raw = re.split(
        r"\s+(?:size|gross\s*weight|gross|net\s*(?:weight|wet|wight)|weight)\s*[:：]?",
        raw,
        maxsplit=1,
        flags=re.I,
    )[0]
    raw = re.sub(r"\s+\+\s+", "+", raw)
    parts = raw.split("/")
    return clean_color_list(parts)


def safe_name(value):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("_")[:80] or "product"


def image_ext(data):
    if data.startswith(b"\x89PNG"):
        return ".png"
    if data.startswith(b"\xff\xd8"):
        return ".jpg"
    if data.startswith(b"GIF"):
        return ".gif"
    if data.startswith(b"RIFF") and b"WEBP" in data[:16]:
        return ".webp"
    return ".png"


def row_dict(headers, values):
    return {headers[index]: values[index] for index in range(min(len(headers), len(values))) if headers[index]}


def extract_row_images(workbook):
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    row_images = {}
    for ws in workbook.worksheets:
        by_row = {}
        for index, img in enumerate(getattr(ws, "_images", [])):
            try:
                row = img.anchor._from.row + 1
                col = img.anchor._from.col + 1
            except Exception:
                continue
            by_row.setdefault(row, []).append(
                {
                    "index": index,
                    "col": col,
                    "width": int(getattr(img, "width", 0) or 0),
                    "height": int(getattr(img, "height", 0) or 0),
                    "img": img,
                }
            )

        for row, images in by_row.items():
            images.sort(key=lambda item: (item["col"] != 5, item["col"], item["index"]))
            row_images[(ws.title, row)] = images
    return row_images


def write_images(sku, row_images):
    paths = []
    for index, image in enumerate(row_images, 1):
        data = image["img"]._data()
        file_name = f"{safe_name(sku)}_{index}{image_ext(data)}"
        file_path = IMAGE_DIR / file_name
        file_path.write_bytes(data)
        paths.append(f"./src/assets/products/{file_name}")
    return paths


def product_score(product):
    return (
        1 if product.get("name") and product["name"] != product["sku"] else 0,
        len(product.get("images") or []),
        len(product.get("description") or ""),
        product.get("price") or 0,
    )


def unique_row_id(sku, counts):
    counts[sku] = counts.get(sku, 0) + 1
    if counts[sku] == 1:
        return sku
    return f"{sku}--{counts[sku]}"


def build_catalog(source):
    workbook = openpyxl.load_workbook(source)
    row_images = extract_row_images(workbook)
    products = []
    sku_counts = {}
    next_sort_order = 1

    for ws in workbook.worksheets:
        headers = [clean(cell.value) for cell in ws[1]]
        for row_number in range(2, ws.max_row + 1):
            row = row_dict(headers, [cell.value for cell in ws[row_number]])
            sku = clean(row.get("SKU"))
            if not sku:
                continue

            raw_info = row.get("Product information") or row.get("Informations") or row.get("Information") or ""
            info = clean(raw_info)
            shade_source = clean(row.get("Colors"))
            name = clean(row.get("Item Name")) or sku
            brand = clean(row.get("Brand")) or "UNBRANDED"
            product_id = unique_row_id(sku, sku_counts)
            images = write_images(product_id, row_images.get((ws.title, row_number), []))
            product = {
                "id": product_id,
                "sku": sku,
                "brand": brand,
                "name": name,
                "category": clean(row.get("Category")) or category(name, info),
                "price": round(number(row.get("Price")), 2),
                "stock": int(number(row.get("QTY"))),
                "weight": round(number(row.get("Weight(kg)"), number(row.get("Weight"))), 3),
                "description": info or "Imported beauty item for wholesale and sample orders.",
                "colors": colors(info, shade_source),
                "image": images[0] if images else "",
                "images": images,
                "published": True,
                "sort_order": next_sort_order,
            }

            products.append(product)
            next_sort_order += 1

    return products


def write_catalog_files(products):
    from server import write_static_catalog
    write_static_catalog(products)


def main():
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    products = build_catalog(source)
    write_catalog_files(products)
    print(f"Wrote {len(products)} products with {sum(1 for p in products if p['image'])} main images and {sum(len(p.get('images') or []) for p in products)} total images.")


if __name__ == "__main__":
    main()
